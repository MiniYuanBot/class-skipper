"""Read every page/body element; preserve locations without a claim-ID graph."""

from pathlib import Path

from storage import digest


def read_materials(slides, transcripts):
    units, sources, warnings = [], [], []
    for index, (path, role) in enumerate(
        [(p, "lecture") for p in slides] + [(p, "transcript") for p in transcripts], 1
    ):
        path = Path(path).resolve(strict=True)
        if path.stat().st_size > 150 * 1024 * 1024:
            raise ValueError("Source exceeds 150 MiB.")
        source = {
            "id": f"s{index}",
            "path": str(path),
            "name": path.name,
            "role": role,
            "sha256": digest(path.read_bytes()),
        }
        sources.append(source)
        suffix = path.suffix.lower()
        if suffix == ".pdf":
            import pypdfium2 as pdfium

            document = pdfium.PdfDocument(path)
            try:
                for i in range(len(document)):
                    page = document[i]
                    textpage = page.get_textpage()
                    try:
                        text = textpage.get_text_range().replace("\r\n", "\n").strip()
                    finally:
                        textpage.close()
                        page.close()
                    units.append(
                        {
                            "id": f"s{index}p{i + 1}",
                            "source": source["id"],
                            "location": f"PDF p.{i + 1}",
                            "page": i + 1,
                            "text": text,
                            "role": role,
                        }
                    )
                    if len(text) < 40:
                        warnings.append(
                            f"s{index}p{i + 1}: sparse text; inspect original or enable vision."
                        )
            finally:
                document.close()
        elif suffix == ".docx":
            from docx import Document
            from docx.oxml.ns import qn

            document = Document(path)
            parts, start, size, number = [], 1, 0, 0
            for i, element in enumerate(document.element.body, 1):
                text = " ".join(t.text or "" for t in element.iter(qn("w:t"))).strip()
                if not text:
                    continue
                parts.append(text)
                size += len(text)
                if size >= 3000:
                    number += 1
                    units.append(
                        {
                            "id": f"s{index}b{number}",
                            "source": source["id"],
                            "location": f"DOCX body {start}-{i}",
                            "text": "\n".join(parts),
                            "role": role,
                        }
                    )
                    parts, start, size = [], i + 1, 0
            if parts:
                number += 1
                units.append(
                    {
                        "id": f"s{index}b{number}",
                        "source": source["id"],
                        "location": f"DOCX body {start}-{len(document.element.body)}",
                        "text": "\n".join(parts),
                        "role": role,
                    }
                )
        elif suffix == ".pptx":
            from pptx import Presentation

            presentation = Presentation(path)

            def shape_text(shapes):
                result = []
                for shape in shapes:
                    if hasattr(shape, "shapes"):
                        result.extend(shape_text(shape.shapes))
                    if shape.has_text_frame:
                        result.append(shape.text)
                    if shape.has_table:
                        result.extend(
                            " | ".join(c.text for c in row.cells) for row in shape.table.rows
                        )
                return result

            for i, slide in enumerate(presentation.slides, 1):
                text = "\n".join(shape_text(slide.shapes))
                if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
                    text += "\nSpeaker notes:\n" + slide.notes_slide.notes_text_frame.text
                units.append(
                    {
                        "id": f"s{index}p{i}",
                        "source": source["id"],
                        "location": f"Slide {i}",
                        "page": i,
                        "text": text,
                        "role": role,
                    }
                )
            warnings.append(
                f"{source['id']}: PPTX diagrams are text-only; use a PDF export for figures."
            )
        elif suffix in {".txt", ".md"}:
            text = path.read_text(encoding="utf-8-sig")
            # Split at paragraph boundaries without dropping the final tail.
            parts, size, number = [], 0, 0
            for paragraph in text.splitlines():
                parts.append(paragraph)
                size += len(paragraph)
                if size >= 3000:
                    number += 1
                    units.append(
                        {
                            "id": f"s{index}b{number}",
                            "source": source["id"],
                            "location": f"Text segment {number}",
                            "text": "\n".join(parts),
                            "role": role,
                        }
                    )
                    parts, size = [], 0
            if parts:
                units.append(
                    {
                        "id": f"s{index}b{number + 1}",
                        "source": source["id"],
                        "location": f"Text segment {number + 1}",
                        "text": "\n".join(parts),
                        "role": role,
                    }
                )
        else:
            raise ValueError(
                "Supported sources: PDF, PPTX, DOCX, TXT, MD. Export legacy PPT to PDF."
            )
    if not any(u["text"].strip() for u in units):
        warnings.append("No extractable text. Visual recognition is required for scanned material.")
    return {"sources": sources, "units": units, "warnings": warnings}


def page_count(source):
    import pypdfium2 as pdfium

    if Path(source["path"]).suffix.lower() != ".pdf":
        raise ValueError("Page rendering currently requires a PDF source.")
    document = pdfium.PdfDocument(source["path"])
    try:
        return len(document)
    finally:
        document.close()


def page_images(source, numbers, scale=1.5):
    """Yield (page number, PIL image), opening the PDF once for many pages."""
    import pypdfium2 as pdfium

    if Path(source["path"]).suffix.lower() != ".pdf":
        raise ValueError("Page rendering currently requires a PDF source.")
    document = pdfium.PdfDocument(source["path"])
    try:
        for number in numbers:
            if not 1 <= number <= len(document):
                raise ValueError("PDF page number is outside the document.")
            page = document[number - 1]
            try:
                bitmap = page.render(scale=scale)
                try:
                    yield number, bitmap.to_pil().convert("RGB")
                finally:
                    bitmap.close()
            finally:
                page.close()
    finally:
        document.close()


def render_page(source, page_number, destination, scale=1.5):
    for _, image in page_images(source, [page_number], scale):
        destination.parent.mkdir(parents=True, exist_ok=True)
        image.save(destination, format="PNG")

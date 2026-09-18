"""Control-flow tests use explicit model doubles; parser tests use real local files."""

import json
from pathlib import Path

import httpx
import pytest

from class_skipper.cli import run_batch
from class_skipper.config import DEFAULTS
from class_skipper.engine import clean_markdown, generate
from class_skipper.llm import Client, ModelError
from class_skipper.materials import read_materials, render_page
from class_skipper.storage import managed


class FakeClient:
    calls_seen = []
    fail_review = False
    fail_section = False

    def __init__(self, config, cache, refresh=False):
        self.calls = []

    def ask(self, stage, system, data, **kwargs):
        self.calls.append({"stage": stage, "cache_hit": False, "usage": None})
        self.calls_seen.append((stage, data))
        ids = [u["id"] for u in data.get("materials", [])]
        if stage == "plan":
            assert "TAIL: necessary condition" in json.dumps(data)
            return {
                "title": "测试课程",
                "learning_thread": "先理解对象，再理解边界。",
                "topics": ["定义", "条件"],
                "sections": [
                    {"id": "a", "title": "定义", "source_ids": ids, "figure_ids": []},
                    {"id": "b", "title": "条件", "source_ids": ids, "figure_ids": []},
                ],
            }
        if stage.startswith("write"):
            if self.fail_section and stage.endswith("2"):
                raise ModelError("simulated network failure")
            return {
                "markdown": "### 定义\n\n**定义**：条件不能省略。\n\n"
                "### 常见问题与解答\n\n**问：何时成立？**\n\n满足条件时。",
                "summary": "保持条件",
                "source_ids": ids,
                "uncertainties": [],
            }
        if stage == "review":
            if self.fail_review:
                raise ModelError("simulated timeout")
            return {
                "introduction": "完整导读",
                "synthesis": "条件限定适用范围。",
                "replacements": [
                    {
                        "section_id": "section-2",
                        "markdown": "### 必要条件\n\n不能省略条件。\n\n### 常见问题与解答\n\n"
                        "**问：缺少条件可以吗？**\n\n不可以。",
                        "source_ids": ids,
                    }
                ],
                "additions": [],
                "changes": ["澄清必要条件"],
                "uncertainties": [],
            }
        raise AssertionError(stage)


@pytest.fixture
def material(tmp_path):
    source = tmp_path / "lecture.txt"
    source.write_text(
        "Definition and examples.\n" + "source explanation\n" * 400 + "TAIL: necessary condition"
    )
    FakeClient.calls_seen = []
    FakeClient.fail_review = FakeClient.fail_section = False
    config = DEFAULTS | {
        "output": str(tmp_path / "out"),
        "workspace": str(tmp_path / "workspace"),
        "allow_remote_llm": True,
    }
    return source, config


def test_full_read_plan_section_write_revision(material):
    source, config = material
    result, code = generate(config, "course", "lecture", [source], client_factory=FakeClient)
    assert code == 0
    note = Path(result["notes"]).read_text()
    assert "必要条件" in note and "本讲小结" in note
    assert result["sections"] == 2
    assert result["review_status"] == "completed"
    assert len(FakeClient.calls_seen) == 4  # plan + 2 chapters + one revision
    assert "[^section-2]:" in note
    assert "TAIL" in json.dumps(next(d for s, d in FakeClient.calls_seen if s == "review"))


def test_review_failure_preserves_draft(material):
    source, config = material
    FakeClient.fail_review = True
    result, code = generate(config, "c", "l", [source], client_factory=FakeClient)
    assert code == 0
    assert result["review_status"] == "unavailable"
    assert "当前保留完整初稿" in Path(result["notes"]).read_text()


def test_section_failure_is_visible_and_nonzero(material):
    source, config = material
    FakeClient.fail_section = True
    result, code = generate(config, "c", "l", [source], client_factory=FakeClient)
    assert code == 3 and result["status"] == "partial"
    assert "本节生成暂未完成" in Path(result["notes"]).read_text()


def test_manual_notes_and_extra_files_are_preserved(material):
    source, config = material
    result, _ = generate(config, "c", "l", [source], client_factory=FakeClient)
    note = Path(result["notes"])
    note.write_text(note.read_text() + "\nMy personal understanding.\n")
    result, code = generate(config, "c", "l", [source], client_factory=FakeClient)
    assert code == 5 and "candidate" in result
    assert "My personal understanding" in note.read_text()
    assert Path(result["candidate"]).exists()


def test_batch_continues_and_indexes_only_existing_notes(material, tmp_path):
    source, config = material
    manifest = tmp_path / "course.yaml"
    manifest.write_text(
        "lectures:\n  - id: first\n    slides: [lecture.txt]\n"
        "  - id: missing\n    slides: [absent.pdf]\n"
        "  - id: last\n    slides: [lecture.txt]\n"
    )

    def runner(*args, **kwargs):
        return generate(*args, **kwargs, client_factory=FakeClient)

    result, code = run_batch(config, "c", manifest, runner=runner)
    assert code == 3
    assert [r["exit_code"] for r in result["lectures"]] == [0, 3, 0]
    index = Path(result["index"]).read_text()
    assert "first/notes.md" in index and "last/notes.md" in index
    assert "missing/notes.md" not in index


def test_source_limit_never_truncates(material):
    source, config = material
    config["max_source_chars"] = 20
    with pytest.raises(ValueError, match="No text was truncated"):
        generate(config, "c", "l", [source], client_factory=FakeClient)
    assert not FakeClient.calls_seen


def test_output_paths_and_markup(material, tmp_path):
    source, config = material
    with pytest.raises(ValueError):
        generate(config, "../escape", "l", [source], client_factory=FakeClient)
    link = tmp_path / "link"
    link.symlink_to(tmp_path / "out")
    with pytest.raises(ValueError):
        managed(link, "c")
    text = clean_markdown(
        "# bad\n![remote](https://invalid)\n<script>alert(1)</script>\n"
        "$$\nx < y > z\n$$\n```c\n#include <stdio.h>\n```"
    )
    assert "### bad" in text and "![remote]" not in text and "<script>" not in text
    assert "x < y > z" in text and "#include <stdio.h>" in text


def test_llm_cache_and_permission(tmp_path, monkeypatch):
    monkeypatch.setenv("DEEPSEEK_API_KEY", "synthetic-key")
    monkeypatch.setenv("CLASS_SKIPPER_TEXT_MODEL", "test-model")
    monkeypatch.delenv("CLASS_SKIPPER_DEEPSEEK_BASE_URL", raising=False)
    requests = []

    def respond(request):
        requests.append(request)
        return httpx.Response(
            200,
            json={
                "choices": [
                    {
                        "finish_reason": "stop",
                        "message": {"content": '{"markdown":"body", "source_ids":[]}'},
                    }
                ],
                "usage": {"prompt_tokens": 10, "completion_tokens": 5, "total_tokens": 15},
            },
        )

    config = DEFAULTS | {"allow_remote_llm": True}
    client = Client(config, tmp_path, transport=httpx.MockTransport(respond))
    first = client.ask("write", "instructions", {"text": "complete source"})
    assert client.ask("write", "instructions", {"text": "complete source"}) == first
    assert len(requests) == 1 and client.calls[-1]["cache_hit"]
    assert requests[0].extensions["timeout"]["connect"] == 15
    assert requests[0].extensions["timeout"]["read"] == config["timeout_seconds"]
    client.ask("write", "changed instructions", {"text": "complete source"})
    assert len(requests) == 2
    blocked = Client(DEFAULTS, tmp_path, transport=httpx.MockTransport(respond))
    with pytest.raises(ModelError, match="permission"):
        blocked.ask("write", "instructions", {})
    assert "synthetic-key" not in "".join(p.read_text() for p in tmp_path.glob("*.json"))


def small_pdf(path):
    streams = [
        b"BT /F1 14 Tf 20 100 Td (First page definition) Tj ET",
        b"BT /F1 14 Tf 20 100 Td (Final necessary condition) Tj ET",
    ]
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R 5 0 R] /Count 2 >>",
    ]
    for page, stream in zip((3, 5), streams):
        objects.append(
            (
                "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 300 200] "
                f"/Resources << /Font << /F1 7 0 R >> >> /Contents {page + 1} 0 R >>"
            ).encode()
        )
        objects.append(f"<< /Length {len(stream)} >>\nstream\n".encode() + stream + b"\nendstream")
    objects.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    raw, offsets = b"%PDF-1.4\n", [0]
    for index, obj in enumerate(objects, 1):
        offsets.append(len(raw))
        raw += f"{index} 0 obj\n".encode() + obj + b"\nendobj\n"
    start = len(raw)
    raw += f"xref\n0 {len(offsets)}\n0000000000 65535 f \n".encode()
    raw += b"".join(f"{i:010d} 00000 n \n".encode() for i in offsets[1:])
    raw += f"trailer\n<< /Size {len(offsets)} /Root 1 0 R >>\nstartxref\n{start}\n%%EOF".encode()
    path.write_bytes(raw)


def test_real_pdf_docx_and_page_render(tmp_path):
    from docx import Document
    from PIL import Image

    pdf, docx = tmp_path / "source.pdf", tmp_path / "speech.docx"
    small_pdf(pdf)
    doc = Document()
    doc.add_paragraph("Instructor explanation.")
    doc.add_table(rows=1, cols=1).cell(0, 0).text = "Table boundary condition."
    doc.add_paragraph("FINAL instructor-only condition.")
    doc.save(docx)
    material = read_materials([pdf], [docx])
    assert len([u for u in material["units"] if u.get("page")]) == 2
    text = "\n".join(u["text"] for u in material["units"])
    assert "Final necessary condition" in text
    assert "Table boundary condition" in text and "FINAL instructor-only" in text
    destination = tmp_path / "figure.png"
    render_page(material["sources"][0], 2, destination)
    with Image.open(destination) as image:
        assert image.size == (450, 300)


def test_json_repair_receives_actual_answer_and_precise_error(tmp_path, monkeypatch):
    from class_skipper.engine import validate_section

    monkeypatch.setenv("DEEPSEEK_API_KEY", "synthetic-key")
    monkeypatch.setenv("CLASS_SKIPPER_TEXT_MODEL", "test-model")
    monkeypatch.delenv("CLASS_SKIPPER_DEEPSEEK_BASE_URL", raising=False)
    requests = []

    def respond(request):
        requests.append(json.loads(request.content))
        value = {"markdown": "complete chapter"}
        if len(requests) > 1:
            value["source_ids"] = ["s1p1"]
        return httpx.Response(
            200,
            json={
                "choices": [{"finish_reason": "stop", "message": {"content": json.dumps(value)}}]
            },
        )

    client = Client(
        DEFAULTS | {"allow_remote_llm": True}, tmp_path, transport=httpx.MockTransport(respond)
    )
    result = client.ask("write", "instructions", {}, validate=validate_section)
    assert result["source_ids"] == ["s1p1"] and len(requests) == 2
    assert requests[1]["messages"][-2]["role"] == "assistant"
    assert "Missing chapter sources" in requests[1]["messages"][-1]["content"]


def test_optional_vision_reads_real_render_and_keeps_relative_asset(tmp_path):
    pdf = tmp_path / "source.pdf"
    small_pdf(pdf)
    images = []

    class VisionDouble(FakeClient):
        def ask(self, stage, system, data, **kwargs):
            if stage.startswith("vision:"):
                assert kwargs["vision"] is True
                images.append(kwargs["image"])
                if stage.startswith("vision:locate"):
                    return {
                        "regions": [
                            {
                                "bbox": [0.05, 0.1, 0.8, 0.8],
                                "caption": "Mechanism",
                                "explanation": "Flow",
                                "section_id": "section-1",
                                "score": 5,
                            }
                        ]
                    }
                return {
                    "usable": True,
                    "caption": "Mechanism",
                    "explanation": "Flow",
                    "reason": "complete",
                }
            if stage == "plan":
                return {
                    "title": "Example",
                    "sections": [
                        {"title": "Topic", "source_ids": ["s1p2"], "figure_ids": ["s1p2"]}
                    ],
                }
            if stage.startswith("write"):
                return {
                    "markdown": "### Concept\n\nExplanation.\n\n"
                    + "\n\n".join(v["marker"] for v in data["visual_readings"].values())
                    + "\n\n### Question\n\nWhy?",
                    "source_ids": ["s1p2"],
                }
            raise AssertionError(stage)

    config = DEFAULTS | {
        "output": str(tmp_path / "output"),
        "workspace": str(tmp_path / "workspace"),
        "allow_remote_llm": True,
        "vision": True,
        "review": False,
    }
    result, code = generate(config, "c", "l", [pdf], client_factory=VisionDouble)
    assert code == 0 and len(images) == 4
    assert all(image.startswith(b"\x89PNG") for image in images)
    path = Path(result["notes"])
    assert "assets/s1p2-f1.png" in path.read_text()
    assert (path.parent / "assets/s1p2-f1.png").exists()


def test_template_table_math_and_html_cleanup():
    value = clean_markdown(
        "| A | B |\n|---|:---:|\n| 1 | 2 |\n"
        "$$x < y > z$$\n\nInline $x < y > z$ and `<stdio.h>`.\n"
        '<svg onload="bad()"><path /></svg>'
    )
    assert "| --- | :---: |" in value
    assert "$$\nx < y > z\n$$" in value
    assert "$x < y > z$" in value and "`<stdio.h>`" in value
    assert "<svg" not in value


def test_output_contains_only_final_files(material):
    source, config = material
    generate(config, "c", "l", [source], client_factory=FakeClient)
    result, code = generate(config, "c", "l", [source], client_factory=FakeClient)
    assert code == 0
    assert [p.name for p in Path(result["notes"]).parent.iterdir()] == ["notes.md"]
    assert list(Path(config["workspace"]).glob("runs/c/l/*/review.json"))
    assert list(Path(config["workspace"]).glob("backups/c/*/notes.md"))


def test_crop_coordinates_and_invalid_boxes(tmp_path):
    from PIL import Image

    from class_skipper.vision import crop_region, valid_bbox

    pdf = tmp_path / "source.pdf"
    small_pdf(pdf)
    destination = tmp_path / "crop.png"
    crop_region({"path": str(pdf)}, 1, [0, 0, 0.5, 0.5], destination)
    with Image.open(destination) as image:
        assert image.size == (450, 300)
    for box in ([0, 0, 1, 2], [1, 0, 0, 1], [0, 0, float("nan"), 1]):
        with pytest.raises(ValueError):
            valid_bbox(box)


def test_vision_rejects_crop_without_whole_page_fallback(tmp_path):
    from class_skipper.vision import collect_figures

    pdf = tmp_path / "source.pdf"
    small_pdf(pdf)
    material = read_materials([pdf], [])

    class RejectingVision:
        def ask(self, stage, system, data, **kwargs):
            if stage.startswith("vision:locate"):
                return {
                    "regions": [
                        {
                            "bbox": [0.1, 0.1, 0.8, 0.8],
                            "caption": "Candidate",
                            "explanation": "Possible flow",
                            "section_id": "section-1",
                            "score": 4,
                        }
                    ]
                }
            return {"usable": False, "caption": "", "explanation": "", "reason": "cut labels"}

    figures, readings, warnings = collect_figures(
        RejectingVision(),
        DEFAULTS,
        tmp_path / "run",
        {u["id"]: u for u in material["units"]},
        {s["id"]: s for s in material["sources"]},
        {"sections": [{"id": "section-1", "title": "Mechanism"}]},
        lambda _: None,
    )
    assert not figures and not readings and len(warnings) == 2


def test_output_workspace_cannot_overlap(material):
    source, config = material
    config["workspace"] = config["output"] + "/intermediate"
    with pytest.raises(ValueError, match="non-nested"):
        generate(config, "c", "l", [source], client_factory=FakeClient)


def test_verified_image_is_inline_and_unknown_markers_are_removed():
    from class_skipper.engine import render

    note = render(
        "Title",
        "Intro",
        [
            {
                "id": "section-1",
                "title": "Topic",
                "markdown": "### Mechanism\n\nExplain arrows.\n\n[[figure:good]]\n\n"
                "[[figure:bad]]\n\n[[figure:good]]\n\n### Question\n\nWhy?",
                "source_ids": [],
                "figure_ids": ["good"],
            }
        ],
        "",
        {},
        {},
        {"good": {"caption": "Actual mechanism"}},
        [],
    )
    assert note.count("assets/good.png") == 1
    assert note.index("Explain arrows") < note.index("assets/good.png") < note.index("### Question")
    assert "[[figure:" not in note


def test_kimi_structural_retry_preserves_reasoning(tmp_path, monkeypatch):
    monkeypatch.setenv("MOONSHOT_API_KEY", "synthetic-key")
    monkeypatch.setenv("CLASS_SKIPPER_VISION_MODEL", "kimi-k3")
    monkeypatch.setenv("CLASS_SKIPPER_KIMI_BASE_URL", "https://api.moonshot.cn/v1")
    requests = []

    def respond(request):
        requests.append(json.loads(request.content))
        answer = "invalid JSON" if len(requests) == 1 else '{"regions":[]}'
        return httpx.Response(
            200,
            json={
                "choices": [
                    {
                        "finish_reason": "stop",
                        "message": {"content": answer, "reasoning_content": "synthetic reasoning"},
                    }
                ]
            },
        )

    client = Client(
        DEFAULTS | {"allow_remote_llm": True}, tmp_path, transport=httpx.MockTransport(respond)
    )
    assert client.ask("vision:test", "instructions", {}, vision=True) == {"regions": []}
    assert requests[0]["reasoning_effort"] == "low" and "thinking" not in requests[0]
    assert requests[1]["messages"][-2]["reasoning_content"] == "synthetic reasoning"

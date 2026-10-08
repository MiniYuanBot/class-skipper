"""Offline skill helper checks with real tiny parsers; no model or course/API run."""

import json
import shutil
import struct
import subprocess
import sys
import tempfile
import unittest
import zlib
from pathlib import Path
from urllib.parse import unquote

SCRIPTS = Path(__file__).resolve().parents[1] / "skills" / "class-skipper" / "scripts"


def small_pdf(path):
    """Write a real one-page PDF with a valid xref, using only the standard library."""
    stream = b"BT /F1 12 Tf 20 150 Td (Local PDF fixture tail condition) Tj ET"
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 200 200] "
        b"/Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
    ]
    data, offsets = b"%PDF-1.4\n", [0]
    for number, body in enumerate(objects, 1):
        offsets.append(len(data))
        data += str(number).encode() + b" 0 obj\n" + body + b"\nendobj\n"
    xref = len(data)
    data += b"xref\n0 6\n0000000000 65535 f \n"
    data += b"".join(f"{offset:010} 00000 n \n".encode() for offset in offsets[1:])
    data += b"trailer\n<< /Size 6 /Root 1 0 R >>\nstartxref\n" + str(xref).encode() + b"\n%%EOF\n"
    path.write_bytes(data)


def note_name(section):
    return section["title"].split()[0] + "-" + section.get("slug", section["id"]) + ".md"


def small_png(path):
    """A real one-pixel PNG fixture; no model-generated or downloaded image."""

    def chunk(kind, payload):
        return (
            struct.pack(">I", len(payload))
            + kind
            + payload
            + struct.pack(">I", zlib.crc32(kind + payload))
        )

    path.write_bytes(
        b"\x89PNG\r\n\x1a\n"
        + chunk(b"IHDR", struct.pack(">2I5B", 1, 1, 8, 2, 0, 0, 0))
        + chunk(b"IDAT", zlib.compress(b"\x00\xff\xff\xff"))
        + chunk(b"IEND", b"")
    )


def model_section(section_id="concept", title="01 性能指标", slug="performance", body="知识正文。"):
    """Labeled model double for publication/check tests, never a course run."""
    return {
        "id": section_id,
        "title": title,
        "slug": slug,
        "summary": "说明指标含义",
        "markdown": "### 定义\n\n" + body + "\n\n> [!question]- 自测：问题？\n> 答案。",
        "source_ids": ["s1b1"],
    }


class OfflineSkillTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.base = Path(self.temp.name)
        self.vault = self.base / "test-vault"
        (self.vault / ".obsidian").mkdir(parents=True)
        self.root = self.vault / "computer-organization-and-architecture"
        self.text = self.base / "lecture.txt"
        self.text.write_text(
            "完整文本\n" + "长段落\n" * 1200 + "TAIL 必要条件", encoding="utf-8-sig"
        )

    def cli(self, *args, expected=0):
        result = subprocess.run(
            [sys.executable, "-X", "utf8", str(SCRIPTS / "local.py"), *map(str, args)],
            capture_output=True,
            text=True,
            encoding="utf-8",
            check=False,
        )
        self.assertEqual(result.returncode, expected, result.stdout + result.stderr)
        return json.loads(result.stdout)

    def prepare(self, *extra, lecture="01"):
        return self.cli(
            "prepare",
            "--root",
            self.root,
            "--course",
            "demo",
            "--lecture",
            lecture,
            "--title",
            "课程标题",
            "--slides",
            self.text,
            *extra,
        )

    def prepare_legacy(self):
        prepared = self.prepare()
        run = Path(prepared["run"])
        metadata = json.loads((run / "run.json").read_text(encoding="utf-8"))
        metadata.pop("layout")
        metadata["lecture"] = "01"
        old_run = self.root / "workspace" / "demo" / "01" / run.name
        old_run.parent.mkdir(parents=True)
        shutil.move(str(run), old_run)
        (old_run / "run.json").write_text(json.dumps(metadata), encoding="utf-8")
        (self.root / "workspace" / "course.json").unlink()
        return str(old_run)

    def write_json(self, filename, value):
        path = self.base / filename
        path.write_text(json.dumps(value, ensure_ascii=False), encoding="utf-8")
        return path

    def document(self, sections=None):
        return self.write_json(
            "document.json",
            {
                "schema_version": 1,
                "title": "完整讲义",
                "introduction": "先看完整材料。",
                "synthesis": "必要条件不能省略。",
                "uncertainties": [],
                "sections": sections
                or [
                    {
                        "id": "concept",
                        "title": "概念",
                        "markdown": (
                            "### 定义\r\n\r\n#### 细节\r\n\r\n```python\r\n"
                            "### literal comment\r\n```\r\n\r\n$$\r\n"
                            "### literal math\r\n$$\r\n\r\n![原图](assets/figure.png)"
                        ),
                        "source_ids": ["s1b1"],
                    },
                    {
                        "id": "condition",
                        "title": "条件",
                        "markdown": "### 必要条件\n\n保留原材料最后一段。",
                        "source_ids": ["s1b2"],
                    },
                ],
            },
        )

    def publish_document(self):
        prepared = self.prepare("--options", '{"course_name":"演示课程"}')
        asset = self.base / "figure.png"
        asset.write_bytes(b"referenced visual")
        document = self.document()
        command = ("publish", "--run", prepared["run"], "--document", document, "--asset", asset)
        return prepared, self.cli(*command), command

    def assert_local_links(self, root):
        import re

        for page in root.rglob("*.md"):
            for reference in re.findall(r"\]\(([^()]+)\)", page.read_text(encoding="utf-8")):
                if reference.startswith("https://"):
                    continue
                self.assertTrue(
                    (page.parent / reference).is_file(), f"Broken link in {page}: {reference}"
                )

    def test_structured_publication_navigation_sources_and_intermediate_layout(self):
        prepared, result, _ = self.publish_document()
        run = Path(prepared["run"])
        for directory in (
            "requests",
            "chapters",
            "visuals/pages",
            "visuals/crops",
            "draft",
            "revision",
            "final",
            "cache",
        ):
            self.assertTrue((run / directory).is_dir())
        self.assertTrue((self.root / "input").is_dir())
        for filename in ("run.json", "materials.json", "status.json"):
            self.assertEqual(
                json.loads((run / filename).read_text(encoding="utf-8"))["schema_version"], 1
            )
        status = run / "status.json"
        status.write_text('{"schema_version":1,"manual":"keep"}', encoding="utf-8")
        self.prepare("--options", '{"course_name":"演示课程"}')
        self.assertIn("manual", status.read_text(encoding="utf-8"))
        output = self.root / "output"
        self.assert_local_links(output)
        self.assertEqual(result["chapters"], 2)
        chapter = output / "L01" / "chapters" / "concept.md"
        body = chapter.read_text(encoding="utf-8")
        self.assertIn('type: "course-note"', body)
        self.assertIn('section: "concept"', body)
        self.assertIn("> [!info]- 来源\n> - lecture.txt：Text segment 1", body)
        self.assertIn("../assets/figure.png", body)
        self.assertIn("\n## 定义\n", body)
        self.assertIn("\n### 细节\n", body)
        self.assertIn("\n### literal comment\n", body)
        self.assertIn("\n### literal math\n", body)
        self.assertNotIn(b"\r", chapter.read_bytes())
        self.assertFalse(list(output.rglob("*.json")))
        self.assertIn("演示课程", Path(result["index"]).read_text(encoding="utf-8"))

    def test_structured_root_and_chapter_manual_changes_block_whole_publication(self):
        _, result, command = self.publish_document()
        root_index = Path(result["index"])
        root_index.write_text("manual library navigation", encoding="utf-8")
        conflict = self.cli(*command, expected=5)
        self.assertIn("index.md", conflict["conflicts"])
        self.assertEqual(root_index.read_text(encoding="utf-8"), "manual library navigation")
        root_index.write_bytes((Path(conflict["candidate"]) / "index.md").read_bytes())
        chapter = self.root / "output" / "L01" / "chapters" / "concept.md"
        chapter.write_text("manual chapter correction", encoding="utf-8")
        conflict = self.cli(*command, expected=5)
        self.assertIn("L01/chapters/concept.md", conflict["conflicts"])
        self.assertEqual(chapter.read_text(encoding="utf-8"), "manual chapter correction")

    def test_structured_export_preserves_full_navigation_and_manual_edits(self):
        self.publish_document()
        command = (
            "export",
            "--root",
            self.root,
            "--course",
            "demo",
            "--vault",
            self.base / "vault",
            "--course-name",
            "课程",
        )
        result = self.cli(*command)
        target = Path(result["target"])
        self.assert_local_links(target)
        self.assertEqual({p.name for p in target.iterdir()}, {"input", "output", "workspace"})
        self.assertEqual(list((target / "input").iterdir()), [])
        self.assertEqual(list((target / "workspace").iterdir()), [])
        self.assertFalse((target / ".obsidian").exists())
        chapter = target / "output" / "L01" / "chapters" / "concept.md"
        self.assertTrue(chapter.is_file())
        chapter.write_text("edited in Obsidian", encoding="utf-8")
        self.cli(*command, expected=5)
        self.assertEqual(chapter.read_text(encoding="utf-8"), "edited in Obsidian")

    def test_legacy_structured_runs_retain_original_paths(self):
        run = self.prepare_legacy()
        document = self.document(
            [
                {
                    "id": "concept",
                    "title": "概念",
                    "markdown": "### 定义\n\n知识正文。",
                    "source_ids": ["s1b1"],
                }
            ]
        )
        result = self.cli("publish", "--run", run, "--document", document)
        output = self.root / "output"
        self.assertEqual(Path(result["lecture_index"]), output / "demo/01/index.md")
        self.assert_local_links(output)
        original = Path(result["lecture_index"]).read_bytes()
        self.cli(
            "prepare",
            "--root",
            self.root,
            "--course",
            "demo",
            "--lecture",
            "L02",
            "--title",
            "L02 新讲",
            "--slides",
            self.text,
            expected=2,
        )
        self.assertFalse((self.root / "workspace/course.json").exists())
        self.assertEqual(Path(result["lecture_index"]).read_bytes(), original)
        exported = self.cli(
            "export",
            "--root",
            self.root,
            "--course",
            "demo",
            "--vault",
            self.base / "export-vault",
            "--course-name",
            "旧课程",
        )
        target = Path(exported["target"])
        self.assertTrue((target / "demo/01/chapters/concept.md").is_file())
        self.assert_local_links(target)

    def test_structured_invalid_ids_and_removes_only_unchanged_stale_chapters(self):
        prepared, _, command = self.publish_document()
        document = self.base / "document.json"
        body = json.loads(document.read_text(encoding="utf-8"))
        body["sections"][0]["source_ids"] = ["unknown"]
        document.write_text(json.dumps(body), encoding="utf-8")
        self.cli(*command, expected=2)
        body["sections"][0]["source_ids"] = ["s1b1"]
        body["schema_version"] = 2
        document.write_text(json.dumps(body), encoding="utf-8")
        self.cli(*command, expected=2)
        body["schema_version"] = 1
        body["sections"][1]["id"] = body["sections"][0]["id"]
        document.write_text(json.dumps(body), encoding="utf-8")
        self.cli(*command, expected=2)
        body["sections"][1]["id"] = "condition"
        body["sections"] = [body["sections"][1]]
        document.write_text(json.dumps(body), encoding="utf-8")
        lecture = self.root / "output" / "L01"
        stale = lecture / "chapters" / "concept.md"
        original = stale.read_bytes()
        stale.write_text("manual correction in retired chapter", encoding="utf-8")
        self.cli(*command, expected=5)
        self.assertEqual(stale.read_text(encoding="utf-8"), "manual correction in retired chapter")
        self.assertTrue((lecture / "assets" / "figure.png").is_file())
        stale.write_bytes(original)
        self.cli(*command)
        self.assertFalse((lecture / "chapters" / "concept.md").exists())
        self.assertFalse((lecture / "assets" / "figure.png").exists())
        self.assertTrue((lecture / "chapters" / "condition.md").is_file())
        self.assertTrue(Path(prepared["run"]).is_dir())

    def test_structured_course_navigation_uses_publication_order(self):
        _, result, _ = self.publish_document()
        asset = self.base / "figure.png"
        for lecture in ("l2", "l10"):
            prepared = self.prepare(lecture=lecture)
            self.cli(
                "publish",
                "--run",
                prepared["run"],
                "--document",
                self.base / "document.json",
                "--asset",
                asset,
            )
        index = Path(result["course_index"]).read_text(encoding="utf-8")
        self.assertLess(index.index("](L02/index.md)"), index.index("](L10/index.md)"))

    def test_chapter_neighbors_follow_final_order_and_survive_export(self):
        prepared = self.prepare()
        sections = [
            {
                "id": section_id,
                "title": title,
                "markdown": "### 定义\n\n知识正文。",
                "source_ids": ["s1b1"],
            }
            for section_id, title in (
                ("section-9", "01 延迟"),
                ("added-topic", "02 流水线"),
                ("section-1", "03 吞吐率"),
            )
        ]
        document = self.document(sections)
        self.cli("publish", "--run", prepared["run"], "--document", document)
        output = self.root / "output"
        self.assert_local_links(output)
        for position, section in enumerate(sections):
            body = (output / "L01/chapters" / note_name(section)).read_text(encoding="utf-8")
            self.assertEqual(body.count("上一节："), 2 if position > 0 else 0)
            self.assertEqual(body.count("下一节："), 2 if position < len(sections) - 1 else 0)
            for neighbor, label in ((position - 1, "上一节"), (position + 1, "下一节")):
                if 0 <= neighbor < len(sections):
                    target = sections[neighbor]
                    link = f"[{label}：{target['title']}]({note_name(target)})"
                    self.assertEqual(body.count(link), 2)
                    self.assertLess(body.index(link), body.index("## 定义"))
                    self.assertGreater(body.rindex(link), body.index("> [!info]- 来源"))
        exported = self.cli(
            "export",
            "--root",
            self.root,
            "--course",
            "demo",
            "--vault",
            self.base / "vault",
            "--course-name",
            "课程",
        )
        target = Path(exported["target"])
        self.assert_local_links(target)
        for section in sections:
            relative = Path("L01/chapters") / note_name(section)
            self.assertEqual(
                (target / "output" / relative).read_bytes(), (output / relative).read_bytes()
            )

    def test_single_chapter_has_no_neighbor_links(self):
        prepared = self.prepare()
        document = self.document(
            [
                {
                    "id": "only",
                    "title": "01 性能指标",
                    "markdown": "### 定义\n\n知识正文。",
                    "source_ids": ["s1b1"],
                }
            ]
        )
        self.cli("publish", "--run", prepared["run"], "--document", document)
        output = self.root / "output"
        self.assert_local_links(output)
        body = (output / "L01/chapters/01-only.md").read_text(encoding="utf-8")
        self.assertNotIn("上一节：", body)
        self.assertNotIn("下一节：", body)
        self.assertEqual(body.count("[本讲目录](../index.md)"), 2)

    def test_inline_footnotes_merged_sources_and_remote_images(self):
        prepared = self.prepare("--transcript", self.text)
        document = self.document(
            [
                {
                    "id": "section-1",
                    "title": "01 流水线: 吞吐/延迟",
                    "slug": "pipeline-throughput",
                    "summary": "区分两种指标",
                    "markdown": (
                        "### 定义\n\n吞吐率提升。[^s1b2]\n\n"
                        "![五级流水](https://upload.wikimedia.org/pipeline.svg)"
                    ),
                    "source_ids": ["s1b1", "s1b2", "s2b2"],
                }
            ]
        )
        self.cli("publish", "--run", prepared["run"], "--document", document)
        output = self.root / "output"
        self.assert_local_links(output)
        body = (output / "L01/chapters/01-pipeline-throughput.md").read_text(encoding="utf-8")
        self.assertIn("lecture.txt：Text segment 1–2", body)
        self.assertIn("lecture.txt：Text segment 2", body)
        self.assertIn("[^s1b2]: lecture.txt-Text segment 2", body)
        self.assertNotIn("[^s1b1]", body)
        self.assertIn("](https://upload.wikimedia.org/pipeline.svg)", body)
        self.assertIn('aliases: ["01 流水线: 吞吐/延迟"]', body)
        index = (output / "L01/index.md").read_text(encoding="utf-8")
        self.assertIn("(chapters/01-pipeline-throughput.md)：区分两种指标", index)
        broken = json.loads(document.read_text(encoding="utf-8"))
        broken["sections"][0]["slug"] = "流水线"
        document.write_text(json.dumps(broken, ensure_ascii=False), encoding="utf-8")
        self.cli("publish", "--run", prepared["run"], "--document", document, expected=2)
        broken["sections"][0]["slug"] = "pipeline-throughput"
        broken["sections"][0]["markdown"] += "\n\n误引。[^nope]"
        document.write_text(json.dumps(broken, ensure_ascii=False), encoding="utf-8")
        self.cli("publish", "--run", prepared["run"], "--document", document, expected=2)

    def test_format_check_reports_mechanical_issues_and_doctor_runs(self):
        prepared = self.prepare()
        good = {
            "id": "section-1",
            "title": "01 流水线",
            "slug": "pipelining",
            "summary": "重叠执行",
            "markdown": (
                "### 定义\n\n**补充解释**：正文。[^s1b1]\n\n$$\nx\n$$\n\n"
                "```text\n# literal\n```\n\n> [!question]- 自测：问题？\n> 答案。"
            ),
            "source_ids": ["s1b1"],
        }
        bad = {
            "id": "section-2",
            "title": "冒险",
            "markdown": "## 越级\n\n**补充解释：**正文 [^zz]\n\n$$\nx",
            "source_ids": ["s1b2"],
        }
        document = self.write_json(
            "check.json", {"schema_version": 1, "title": "L01 流水线", "sections": [good, bad]}
        )
        result = self.cli("check", "--run", prepared["run"], "--document", document)
        self.assertEqual(result["status"], "issues")
        self.assertEqual({issue["section"] for issue in result["issues"]}, {"section-2"})
        self.assertEqual(len(result["issues"]), 8)
        self.assertIn(self.cli("doctor")["status"], {"ready", "setup_needed"})

    def test_bitfield_image_labels_publish_export_and_preserve_manual_edits(self):
        prepared = self.prepare()
        asset = self.base / "figure.png"
        small_png(asset)
        original_asset = asset.read_bytes()
        caption = r"图：动态功耗近似为 $P \approx \frac{1}{2} C V^2 A f$。"
        remote = '![远程图](https://example.invalid/figure_(1).png "remote title")'
        export_command = (
            "export",
            "--root",
            self.root,
            "--course",
            "demo",
            "--vault",
            self.base / "export-vault",
            "--course-name",
            "课程",
        )
        for image in (
            "![指令字段](assets/figure.png)",
            "![立即数 imm[0:5] 与 Instr[6-0]](assets/figure.png)",
            r'![字段 imm\[0:5\] 与 Instr\[6-0\]](assets/figure.png "位域图")',
        ):
            with self.subTest(image=image):
                document = self.document(
                    [model_section(body=image + "\n\n" + caption + "\n\n" + remote)]
                )
                command = ("publish", "--run", prepared["run"], "--document", document)
                self.cli(*command, expected=2)
                self.cli(*command, "--asset", asset)
                chapter = self.root / "output/L01/chapters/01-performance.md"
                body = chapter.read_text(encoding="utf-8")
                self.assertIn(image.replace("assets/figure.png", "../assets/figure.png"), body)
                self.assertIn(caption, body)
                self.assertIn(remote, body)
                self.assertEqual(
                    (self.root / "output/L01/assets/figure.png").read_bytes(), original_asset
                )
                exported = self.cli(*export_command)
                target = Path(exported["target"]) / "output/L01"
                self.assertEqual((target / "assets/figure.png").read_bytes(), original_asset)
                self.assertEqual(
                    (target / "chapters/01-performance.md").read_bytes(), chapter.read_bytes()
                )

        exported_asset = target / "assets/figure.png"
        exported_asset.write_bytes(b"manual image edit in the vault")
        conflict = self.cli(*export_command, expected=5)
        self.assertIn("output/L01/assets/figure.png", conflict["conflicts"])
        self.assertEqual(exported_asset.read_bytes(), b"manual image edit in the vault")
        published_asset = self.root / "output/L01/assets/figure.png"
        published_asset.write_bytes(b"manual local image correction")
        previous_chapter = chapter.read_bytes()
        conflict = self.cli(*command, "--asset", asset, expected=5)
        self.assertIn("L01/assets/figure.png", conflict["conflicts"])
        self.assertEqual(published_asset.read_bytes(), b"manual local image correction")
        self.assertEqual(chapter.read_bytes(), previous_chapter)
        self.assertEqual(
            (Path(conflict["candidate"]) / "L01/assets/figure.png").read_bytes(), original_asset
        )

    def test_image_examples_in_code_and_comments_are_not_assets(self):
        prepared = self.prepare()
        examples = (
            "`![Power ≈ ½CV²Af](assets/inline.png)`\n\n"
            "```markdown\n![imm[0:5]](assets/example.png)\n[[不存在]]\n```\n\n"
            "<!-- ![Instr[6-0]](assets/comment.png) -->\n\n"
            "$[x](section:missing)$"
        )
        section = model_section(body=examples)
        document = self.document([section])
        data = json.loads(document.read_text(encoding="utf-8"))
        data["title"] = "L01 性能指标"
        document = self.write_json("document.json", data)
        result = self.cli("check", "--run", prepared["run"], "--document", document)
        self.assertEqual(result["status"], "ok")
        self.cli("publish", "--run", prepared["run"], "--document", document)
        body = (self.root / "output/L01/chapters/01-performance.md").read_text(encoding="utf-8")
        self.assertIn(examples, body)
        self.assertFalse((self.root / "output/L01/assets").exists())
        exported = self.cli(
            "export",
            "--root",
            self.root,
            "--course",
            "demo",
            "--vault",
            self.base / "export-vault",
            "--course-name",
            "课程",
        )
        self.assertFalse((Path(exported["target"]) / "output/L01/assets").exists())

    def test_legacy_publication_and_exports_keep_bitfield_images(self):
        asset = self.base / "figure.png"
        small_png(asset)
        image = '![imm[0:5] 与 Instr[6-0]](assets/figure.png "bit fields")'
        caption = r"图：功耗为 $P \approx \frac{1}{2} C V^2 A f$。"
        for structured in (False, True):
            with self.subTest(structured=structured):
                self.root = self.base / ("legacy-structured" if structured else "legacy-note")
                run = self.prepare_legacy()
                if structured:
                    document = self.document([model_section(body=image + "\n\n" + caption)])
                    command = ("publish", "--run", run, "--document", document)
                    published_relative = "demo/01/chapters/concept.md"
                    exported_relative = published_relative
                    exported_image = image.replace("assets/figure.png", "../assets/figure.png")
                    exported_asset = "demo/01/assets/figure.png"
                else:
                    note = self.base / "legacy-note.md"
                    note.write_text("# 内容\n\n" + image + "\n\n" + caption, encoding="utf-8")
                    command = ("publish", "--run", run, "--note", note)
                    published_relative = "demo/01/notes.md"
                    exported_relative = "01.md"
                    exported_image = image.replace("assets/figure.png", "assets/01/figure.png")
                    exported_asset = "assets/01/figure.png"
                self.cli(*command, expected=2)
                self.cli(*command, "--asset", asset)
                self.assertTrue((self.root / "output" / published_relative).is_file())
                result = self.cli(
                    "export",
                    "--root",
                    self.root,
                    "--course",
                    "demo",
                    "--vault",
                    self.base / "export-vault",
                    "--course-name",
                    "structured" if structured else "note",
                )
                target = Path(result["target"])
                body = (target / exported_relative).read_text(encoding="utf-8")
                self.assertIn(exported_image, body)
                self.assertIn(caption, body)
                self.assertEqual((target / exported_asset).read_bytes(), asset.read_bytes())

    def test_math_check_covers_prose_summaries_indexes_and_image_alt(self):
        prepared = self.prepare()
        good = {
            "schema_version": 1,
            "title": "L01 功耗",
            "introduction": "说明功耗来源。",
            "synthesis": r"动态功耗为 $P \approx \frac{1}{2} C V^2 A f$。",
            "uncertainties": [],
            "sections": [
                model_section(
                    body=(
                        "![功耗示意图](https://example.invalid/power.png)\n\n"
                        r"图：$P \approx \frac{1}{2} C V^2 A f$，$x_0$ 为初值。"
                        "\n\n位域 `imm[0:5]` 与 `Instr[6-0]` 保留代码语义。"
                        "\n\n示例文本 `Power ≈ ½CV²Af`。\n\n```text\nx₀ = ½V²\n```"
                    )
                )
            ],
        }
        document = self.write_json("math.json", good)
        command = ("check", "--run", prepared["run"], "--document", document)
        self.assertEqual(self.cli(*command)["status"], "ok")
        for field, value in (
            ("markdown", "正文 Power ≈ ½CV²Af。"),
            ("markdown", "正文初始变量 x₀。"),
            ("markdown", "正文比例 ⅓。"),
            ("markdown", "![公式 $P = C V^2$](https://example.invalid/power.png)"),
            ("summary", "电压为 V²"),
            ("introduction", "比例为 ½"),
            ("synthesis", "初始变量为 $x₀$"),
        ):
            with self.subTest(field=field, value=value):
                candidate = json.loads(json.dumps(good))
                if field in {"markdown", "summary"}:
                    candidate["sections"][0][field] += "\n\n" + value
                else:
                    candidate[field] = value
                self.write_json("math.json", candidate)
                result = self.cli(*command)
                self.assertEqual(result["status"], "issues")
                self.assertTrue(result["issues"])

    def test_section_links_resolve_forward_titles_fragments_and_reordering(self):
        prepared = self.prepare()
        untouched = (
            "`[示例](section:power)`\n\n"
            "```markdown\n[[02 动态功耗|例子]]\n```\n\n"
            "$[x](section:power)$\n\n"
            "[外部](https://example.invalid/section:power)"
        )
        first = model_section(
            body=(
                "[查看 [推导]](section:power#Power)\n\n"
                "[旧标题](02 动态功耗.md#Power)\n\n"
                "[[02 动态功耗#Power|自定义功耗]]\n\n"
                "[[02 动态功耗]]\n\n"
                "[[02 动态功耗#^power|公式块]]\n\n" + untouched
            )
        )
        second = model_section("power", "02 动态功耗", "dynamic-power", "[指标](section:concept)")
        data = {
            "schema_version": 1,
            "title": "L01 性能与功耗",
            "introduction": "[先看功耗](section:power#Power)",
            "synthesis": "[[01 性能指标|回看指标]]",
            "uncertainties": [],
            "sections": [first, second],
        }
        for sections in ([first, second], [second, first]):
            with self.subTest(order=[section["id"] for section in sections]):
                document = self.write_json("links.json", data | {"sections": sections})
                checked = self.cli("check", "--run", prepared["run"], "--document", document)
                self.assertEqual(checked["status"], "ok")
                self.cli("publish", "--run", prepared["run"], "--document", document)
                lecture = self.root / "output/L01"
                body = (lecture / "chapters/01-performance.md").read_text(encoding="utf-8")
                for expected in (
                    "[查看 [推导]](02-dynamic-power.md#Power)",
                    "[旧标题](02-dynamic-power.md#Power)",
                    "[自定义功耗](02-dynamic-power.md#Power)",
                    "[02 动态功耗](02-dynamic-power.md)",
                    "[公式块](02-dynamic-power.md#^power)",
                ):
                    self.assertIn(expected, unquote(body))
                self.assertIn(untouched, body)
                self.assertIn(
                    "[指标](01-performance.md)",
                    (lecture / "chapters/02-dynamic-power.md").read_text(encoding="utf-8"),
                )
                index = (lecture / "index.md").read_text(encoding="utf-8")
                self.assertIn("[先看功耗](chapters/02-dynamic-power.md#Power)", index)
                self.assertIn("[回看指标](chapters/01-performance.md)", index)
        exported = self.cli(
            "export",
            "--root",
            self.root,
            "--course",
            "demo",
            "--vault",
            self.base / "export-vault",
            "--course-name",
            "课程",
        )
        for relative in ("index.md", "chapters/01-performance.md", "chapters/02-dynamic-power.md"):
            self.assertEqual(
                (Path(exported["target"]) / "output/L01" / relative).read_bytes(),
                (lecture / relative).read_bytes(),
            )

    def test_section_ids_resolve_distinct_files_when_title_and_slug_collide(self):
        prepared = self.prepare()
        sections = [
            model_section("first", "01 重名", "same", "[下一项](section:second)"),
            model_section("second", "01 重名", "same", "[上一项](section:first)"),
        ]
        document = self.document(sections)
        self.cli("publish", "--run", prepared["run"], "--document", document)
        chapters = self.root / "output/L01/chapters"
        self.assertEqual(
            {path.name for path in chapters.iterdir()}, {"01-same.md", "01-same-second.md"}
        )
        self.assertIn(
            "[下一项](01-same-second.md)", (chapters / "01-same.md").read_text(encoding="utf-8")
        )
        self.assertIn(
            "[上一项](01-same.md)", (chapters / "01-same-second.md").read_text(encoding="utf-8")
        )

    def test_second_filename_collision_cannot_overwrite_an_existing_chapter(self):
        prepared = self.prepare()
        sections = [
            model_section("one", "01 Same", "x-b", "第一章独有内容。"),
            model_section("two", "01 Same", "x", "第二章独有内容。"),
        ]
        document = self.document(sections)
        command = ("publish", "--run", prepared["run"], "--document", document)
        self.cli(*command)
        output = self.root / "output"
        before = {
            path.relative_to(output): path.read_bytes()
            for path in output.rglob("*")
            if path.is_file()
        }
        sections.append(model_section("b", "01 Same", "x", "第三章不可覆盖第一章。"))
        self.document(sections)
        self.cli(*command, expected=2)
        after = {
            path.relative_to(output): path.read_bytes()
            for path in output.rglob("*")
            if path.is_file()
        }
        self.assertEqual(after, before)

    def test_encoded_hash_in_chapter_title_is_not_the_fragment_separator(self):
        prepared = self.prepare()
        section = model_section(
            "csharp", "01 C#语言", "csharp", "### Heading\n\n[C#](01%20C%23语言#Heading)"
        )
        document = self.document([section])
        data = json.loads(document.read_text(encoding="utf-8"))
        data["title"] = "L01 编程语言"
        self.write_json("document.json", data)
        checked = self.cli("check", "--run", prepared["run"], "--document", document)
        self.assertEqual(checked["status"], "ok")
        self.cli("publish", "--run", prepared["run"], "--document", document)
        body = (self.root / "output/L01/chapters/01-csharp.md").read_text(encoding="utf-8")
        self.assertIn("[C#](01-csharp.md#Heading)", body)

    def test_unknown_and_ambiguous_internal_links_report_before_publication(self):
        prepared = self.prepare()
        for reference, duplicate_title in (
            ("[失效](section:missing)", False),
            ("[[不存在#Heading|失效]]", False),
            ("[失效](missing.md)", False),
            ("[歧义](01 重名)", True),
        ):
            with self.subTest(reference=reference):
                sections = [model_section("first", "01 重名", "first", reference)]
                if duplicate_title:
                    sections.append(model_section("second", "01 重名", "second"))
                document = self.document(sections)
                data = json.loads(document.read_text(encoding="utf-8"))
                data["title"] = "L01 引用"
                self.write_json("document.json", data)
                result = self.cli("check", "--run", prepared["run"], "--document", document)
                self.assertEqual(result["status"], "issues")
                self.assertTrue(any(issue["section"] == "first" for issue in result["issues"]))
                self.cli("publish", "--run", prepared["run"], "--document", document, expected=2)
                self.assertEqual(list((self.root / "output").iterdir()), [])

    def test_existing_cross_lecture_relative_links_keep_their_destinations(self):
        target = self.prepare(lecture="02")
        document = self.document([model_section("other", "02 其他主题", "other")])
        self.cli("publish", "--run", target["run"], "--document", document)
        prepared = self.prepare()
        reference = "[其他讲义](../../L02/chapters/02-other.md#定义)"
        document = self.document([model_section(body=reference)])
        data = json.loads(document.read_text(encoding="utf-8"))
        data["title"] = "L01 引用"
        data["introduction"] = "[其他目录](../L02/index.md)"
        self.write_json("document.json", data)
        checked = self.cli("check", "--run", prepared["run"], "--document", document)
        self.assertEqual(checked["status"], "ok")
        self.cli("publish", "--run", prepared["run"], "--document", document)
        body = (self.root / "output/L01/chapters/01-performance.md").read_text(encoding="utf-8")
        self.assertIn(reference, unquote(body))
        self.assertIn(
            "[其他目录](../L02/index.md)",
            (self.root / "output/L01/index.md").read_text(encoding="utf-8"),
        )

    def test_cross_lecture_receipt_cannot_replace_a_deleted_target_file(self):
        target = self.prepare(lecture="02")
        document = self.document([model_section("other", "02 其他主题", "other")])
        self.cli("publish", "--run", target["run"], "--document", document)
        receipt = self.root / "workspace/publication/published.json"
        previous_receipt = receipt.read_bytes()
        self.assertIn("L02/chapters/02-other.md", json.loads(previous_receipt)["files"])
        (self.root / "output/L02/chapters/02-other.md").unlink()
        prepared = self.prepare()
        document = self.document(
            [model_section(body="[已删除的章节](../../L02/chapters/02-other.md#定义)")]
        )
        data = json.loads(document.read_text(encoding="utf-8"))
        data["title"] = "L01 引用"
        self.write_json("document.json", data)
        checked = self.cli("check", "--run", prepared["run"], "--document", document)
        self.assertEqual(checked["status"], "issues")
        self.assertTrue(any(issue["section"] == "concept" for issue in checked["issues"]))
        self.cli("publish", "--run", prepared["run"], "--document", document, expected=2)
        self.assertFalse((self.root / "output/L01").exists())
        self.assertEqual(receipt.read_bytes(), previous_receipt)

    def test_complete_utf8_materials_identity_and_no_repository_dependency(self):
        first = self.prepare()
        materials = json.loads(Path(first["materials"]).read_text(encoding="utf-8"))
        self.assertEqual(len(materials["sources"]), 1)
        self.assertIn("TAIL 必要条件", materials["units"][-1]["text"])
        self.assertGreater(len(materials["units"]), 1)
        self.assertEqual(first["run"], self.prepare()["run"])
        self.assertNotEqual(first["run"], self.prepare("--options", '{"language":"en"}')["run"])
        self.assertNotEqual(first["run"], self.prepare("--transcript", self.text)["run"])
        self.text.write_text("updated full source", encoding="utf-8")
        self.assertNotEqual(first["run"], self.prepare()["run"])

    def test_cache_invalidates_prompt_options_source_and_upstream_and_archives_refresh(
        self,
    ):
        run = self.prepare()["run"]
        request = self.write_json(
            "request.json",
            {
                "stage": "write:01",
                "instructions": "full prompt",
                "options": {"language": "zh"},
            },
        )
        response = self.write_json("response.json", {"markdown": "labeled model double response"})
        upstream = self.write_json("plan.json", {"sections": ["01"]})
        command = ("--run", run, "--request", request, "--upstream", upstream)
        self.assertEqual(self.cli("cache", "lookup", *command)["status"], "miss")
        stored = self.cli("cache", "store", *command, "--response", response)
        self.assertEqual(
            self.cli("cache", "lookup", *command)["response"]["markdown"],
            "labeled model double response",
        )
        response.write_text('{"markdown":"new complete response"}', encoding="utf-8")
        self.cli("cache", "store", *command, "--response", response, expected=2)
        self.cli("cache", "store", *command, "--response", response, "--refresh")
        history = list((Path(run) / "cache" / "history").glob("*.json"))
        self.assertEqual(len(history), 1)
        self.assertIn("labeled model double", history[0].read_text(encoding="utf-8"))
        original = json.loads(request.read_text(encoding="utf-8"))
        for key, value in (
            ("instructions", "changed full prompt"),
            ("options", {"language": "en"}),
        ):
            request.write_text(json.dumps(original | {key: value}), encoding="utf-8")
            self.assertEqual(self.cli("cache", "lookup", *command)["status"], "miss")
        request.write_text(json.dumps(original), encoding="utf-8")
        upstream.write_text('{"sections":["02"]}', encoding="utf-8")
        self.assertEqual(self.cli("cache", "lookup", *command)["status"], "miss")
        self.text.write_text("new material", encoding="utf-8")
        self.assertEqual(
            self.cli("cache", "lookup", "--run", self.prepare()["run"], "--request", request)[
                "status"
            ],
            "miss",
        )
        self.assertTrue(Path(stored["cache"]).exists())

    def test_publish_assets_note_and_index_manual_protection(self):
        run = self.prepare_legacy()
        note = self.base / "final.md"
        note.write_text("# 内容\n\n![原图](assets/figure.png)\n", encoding="utf-8")
        asset = self.base / "figure.png"
        unused = self.base / "unused.png"
        asset.write_bytes(b"local image bytes")
        unused.write_bytes(b"not referenced")
        command = (
            "publish",
            "--run",
            run,
            "--note",
            note,
            "--asset",
            asset,
            "--asset",
            unused,
        )
        result = self.cli(*command)
        published = Path(result["notes"])
        index = Path(result["index"])
        self.assertFalse((published.parent / "assets" / unused.name).exists())
        published.write_text("manual note", encoding="utf-8")
        result = self.cli(*command, expected=5)
        self.assertEqual(published.read_text(encoding="utf-8"), "manual note")
        self.assertIn("01/notes.md", result["conflicts"])
        self.assertTrue((Path(result["candidate"]) / "01" / "notes.md").is_file())
        published.write_text(note.read_text(encoding="utf-8"), encoding="utf-8")
        index.write_text("manual index", encoding="utf-8")
        self.cli(*command, expected=5)
        self.assertEqual(index.read_text(encoding="utf-8"), "manual index")

    def test_export_preserves_manual_edits_and_rewrites_images(self):
        run = self.prepare_legacy()
        note = self.base / "final.md"
        note.write_text("# Notes\n\n![Figure](assets/figure.png)", encoding="utf-8")
        asset = self.base / "figure.png"
        asset.write_bytes(b"local image")
        self.cli("publish", "--run", run, "--note", note, "--asset", asset)
        vault = self.base / "vault"
        command = (
            "export",
            "--root",
            self.root,
            "--course",
            "demo",
            "--vault",
            vault,
            "--course-name",
            "测试课程",
        )
        result = self.cli(*command)
        exported = Path(result["target"]) / "01.md"
        self.assertIn("assets/01/figure.png", exported.read_text(encoding="utf-8"))
        self.assertIn("](01.md)", Path(result["index"]).read_text(encoding="utf-8"))
        exported.write_text("manual vault note", encoding="utf-8")
        self.cli(*command, expected=5)
        self.assertEqual(exported.read_text(encoding="utf-8"), "manual vault note")

    def test_portable_names_and_lecture_normalization(self):
        first = self.prepare(lecture="l1")
        self.assertEqual(first["run"], self.prepare(lecture="01")["run"])
        output = self.root / "output"
        (output / "l02").mkdir()
        self.cli(
            "prepare",
            "--root",
            self.root,
            "--course",
            "demo",
            "--lecture",
            "L02",
            "--title",
            "Collision",
            "--slides",
            self.text,
            expected=2,
        )
        for course, lecture in (("CON", "L01"), ("demo", "Lesson")):
            self.cli(
                "prepare",
                "--root",
                self.root,
                "--course",
                course,
                "--lecture",
                lecture,
                "--title",
                "Invalid",
                "--slides",
                self.text,
                expected=2,
            )

    def test_course_roots_are_isolated_and_vault_settings_are_untouched(self):
        settings = self.vault / ".obsidian" / "app.json"
        settings.write_text('{"keep":true}', encoding="utf-8")
        prepared, result, _ = self.publish_document()
        self.assertEqual(Path(prepared["run"]).parent, self.root / "workspace/L01")
        self.assertEqual(Path(result["course_index"]), self.root / "output/index.md")
        self.assertEqual({p.name for p in (self.root / "output").iterdir()}, {"index.md", "L01"})
        self.assertFalse((self.root / ".obsidian").exists())
        second = self.vault / "operating-systems"
        other = self.cli(
            "prepare",
            "--root",
            second,
            "--course",
            "os",
            "--lecture",
            "2",
            "--title",
            "L02 进程",
            "--slides",
            self.text,
        )
        asset = self.base / "figure.png"
        self.cli(
            "publish",
            "--run",
            other["run"],
            "--document",
            self.base / "document.json",
            "--asset",
            asset,
        )
        self.assertTrue((second / "output/L02/index.md").is_file())
        self.assertFalse((second / "output/L01").exists())
        self.assertFalse((self.root / "output/L02").exists())
        self.assertEqual(settings.read_text(encoding="utf-8"), '{"keep":true}')
        self.cli(
            "prepare",
            "--root",
            self.root,
            "--course",
            "another",
            "--lecture",
            "L02",
            "--title",
            "L02 其他课程",
            "--slides",
            self.text,
            expected=2,
        )
        self.cli(
            "prepare",
            "--root",
            self.vault,
            "--course",
            "demo",
            "--lecture",
            "L01",
            "--title",
            "L01 错误根目录",
            "--slides",
            self.text,
            expected=2,
        )
        self.assertFalse((self.vault / "output").exists())

    def test_new_runs_reject_legacy_note_publication_and_keep_sources(self):
        prepared = self.prepare()
        note = self.base / "final.md"
        note.write_text("# 内容", encoding="utf-8")
        self.cli("publish", "--run", prepared["run"], "--note", note, expected=2)
        self.assertTrue(self.text.is_file())
        self.assertEqual(list((self.root / "output").iterdir()), [])

    def test_real_docx_and_pptx_parser_fixtures(self):
        from docx import Document
        from pptx import Presentation
        from pptx.util import Inches

        doc = Document()
        doc.add_paragraph("DOCX first body")
        doc.add_table(rows=1, cols=1).cell(0, 0).text = "DOCX table condition"
        doc.add_paragraph("DOCX final tail")
        docx = self.base / "fixture.docx"
        doc.save(docx)
        deck = Presentation()
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        slide.shapes.add_textbox(
            Inches(1), Inches(1), Inches(3), Inches(1)
        ).text = "PPTX slide concept"
        slide.notes_slide.notes_text_frame.text = "PPTX speaker final condition"
        pptx = self.base / "fixture.pptx"
        deck.save(pptx)
        result = self.prepare("--slides", pptx, "--transcript", docx)
        materials = json.loads(Path(result["materials"]).read_text(encoding="utf-8"))
        text = "\n".join(unit["text"] for unit in materials["units"])
        for expected in (
            "DOCX table condition",
            "DOCX final tail",
            "PPTX speaker final condition",
        ):
            self.assertIn(expected, text)
        self.assertIn("Slide 1", [u["location"] for u in materials["units"]])
        self.assertTrue(any(u["location"].startswith("DOCX body") for u in materials["units"]))

    def test_real_pdf_parser_render_crop_and_changed_source(self):
        from PIL import Image

        pdf = self.base / "fixture.pdf"
        small_pdf(pdf)
        result = self.prepare("--slides", pdf)
        materials = json.loads(Path(result["materials"]).read_text(encoding="utf-8"))
        unit = next(u for u in materials["units"] if u["source"] == "s2")
        self.assertIn("tail condition", unit["text"])
        self.assertEqual(unit["location"], "PDF p.1")
        output = self.base / "figure.png"
        self.cli(
            "render",
            "--run",
            result["run"],
            "--source",
            "s2",
            "--page",
            "1",
            "--crop",
            "0,0,50,60",
            "--output",
            output,
        )
        with Image.open(output) as image:
            self.assertEqual(image.size, (50, 60))
        self.cli(
            "render",
            "--run",
            result["run"],
            "--source",
            "s2",
            "--page",
            "1",
            "--scale",
            "1",
            "--crop",
            "0.5,0,0.5,0.25",
            "--output",
            output,
        )
        with Image.open(output) as image:
            self.assertEqual(image.size, (100, 50))
        sheets = self.cli(
            "sheet", "--run", result["run"], "--source", "s2", "--output-dir", self.base / "sheets"
        )["sheets"]
        self.assertEqual(sheets[0]["pages"], [1])
        with Image.open(sheets[0]["path"]) as image:
            self.assertGreater(image.width, 360)
        self.cli(
            "sheet",
            "--run",
            result["run"],
            "--source",
            "s2",
            "--pages",
            "1-3",
            "--output-dir",
            self.base / "sheets",
            expected=2,
        )
        previous = output.read_bytes()
        self.cli(
            "render",
            "--run",
            result["run"],
            "--source",
            "s2",
            "--page",
            "1",
            "--crop",
            "0,0,9999,9999",
            "--output",
            output,
            expected=2,
        )
        self.assertEqual(output.read_bytes(), previous)
        pdf.write_bytes(pdf.read_bytes() + b"\n% changed source\n")
        self.cli(
            "render",
            "--run",
            result["run"],
            "--source",
            "s2",
            "--page",
            "1",
            "--output",
            output,
            expected=2,
        )


if __name__ == "__main__":
    unittest.main()
